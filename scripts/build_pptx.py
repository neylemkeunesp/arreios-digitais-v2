"""
Regenera os 2 PPTX (light + dark) a partir dos MDs.
Uso: .venv/bin/python scripts/build_pptx.py
"""
import re, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def regen(mdprefix, is_dark):
    src = f"{mdprefix}.md"
    dst = f"{mdprefix}.pptx"
    with open(src) as f:
        text = f.read()
    text = re.sub(r"^---.*?---\n", "", text, count=1, flags=re.DOTALL)
    text = re.sub(r"<style>.*?</style>", "", text, flags=re.DOTALL)
    slides = re.split(r"^---\s*$", text, flags=re.MULTILINE)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    BLANK = prs.slide_layouts[6]

    if is_dark:
        GRAY = RGBColor(0x6c, 0x70, 0x86)
        TITLE_C = RGBColor(0x89, 0xb4, 0xfa)
        TEXT_C = RGBColor(0xcd, 0xd6, 0xf4)
        header = "Arreios Digitais Autoadaptativos e Multimodais · v2 · 🌙 dark"
    else:
        GRAY = RGBColor(0x55, 0x55, 0x55)
        TITLE_C = RGBColor(0x00, 0x30, 0x87)
        TEXT_C = RGBColor(0x22, 0x22, 0x22)
        header = "Arreios Digitais Autoadaptativos e Multimodais · v2"

    def clean(s):
        s = re.sub(r"<[^>]+>", "", s)
        s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)
        s = re.sub(r"\*([^*]+)\*", r"\1", s)
        s = re.sub(r"`([^`]+)`", r"\1", s)
        return s.strip()

    def add_text(slide, text, left, top, width, height, size=14, bold=False, color=None, italic=False):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        lines = text if isinstance(text, list) else text.split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = clean(line)
            for run in p.runs:
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.italic = italic
                if color is not None:
                    run.font.color.rgb = color
        return box

    for i, raw in enumerate(slides):
        raw = raw.strip()
        if not raw:
            continue
        slide = prs.slides.add_slide(BLANK)
        bg = "#1e1e2e" if is_dark else "#ffffff"
        fg = None
        paginate = True
        footer = "Ney Lemke · UNESP · CTInf · 2026"
        body_lines = []
        for line in raw.split("\n"):
            if line.strip().startswith("<!-- _"):
                m = re.search(r"backgroundColor:\s*([#\w]+)", line)
                if m: bg = m.group(1)
                m = re.search(r"color:\s*([#\w]+)", line)
                if m: fg = m.group(1)
                if "_paginate: false" in line: paginate = False
                if "_header: ''" in line: header = ""
                if "_footer: ''" in line: footer = ""
            else:
                body_lines.append(line)
        fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        fill.line.fill.background()
        c = bg.lstrip("#")
        fill.fill.solid()
        fill.fill.fore_color.rgb = RGBColor(int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))
        text_color = RGBColor(int(fg[1:3], 16), int(fg[3:5], 16), int(fg[5:7], 16)) if fg and fg.startswith("#") else None
        if header:
            add_text(slide, [header], Inches(0.3), Inches(0.15), Inches(12.7), Inches(0.3), size=10, color=GRAY, italic=True)
        if footer and paginate:
            add_text(slide, [f"{footer}    ·    {i+1}"], Inches(0.3), Inches(7.15), Inches(12.7), Inches(0.3), size=10, color=GRAY, italic=True)
        body = "\n".join(body_lines).strip()
        if not body:
            continue
        lines = body.split("\n")
        title = ""
        rest_start = 0
        if lines and lines[0].startswith("# "):
            title = lines[0][2:].strip()
            rest_start = 1
        elif lines and lines[0].startswith("## "):
            title = lines[0][3:].strip()
            rest_start = 1
        if title:
            add_text(slide, [title], Inches(0.5), Inches(0.6), Inches(12.3), Inches(1.0), size=32, bold=True, color=TITLE_C)
        rest = "\n".join(lines[rest_start:]).strip()
        if rest:
            top = Inches(1.7) if title else Inches(0.6)
            add_text(slide, rest.split("\n"), Inches(0.5), top, Inches(12.3), Inches(5.3), size=14, color=text_color or TEXT_C)
    prs.save(dst)
    print(f"  {mdprefix}: {len(prs.slides)} slides · {os.path.getsize(dst):,} bytes")

regen("arreios-digitais-v2", False)
regen("arreios-digitais-v2-dark", True)
